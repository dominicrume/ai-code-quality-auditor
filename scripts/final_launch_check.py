#!/usr/bin/env python3
"""Pre-launch audit of everything that leaves this repository.

Checks the submitted document against the data behind it, the data files
against each other, and the artefacts that would be published. Intended to be
run immediately before submission and before deploying the dashboard.

    .venv/bin/python scripts/final_launch_check.py
"""
from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
DOCX = ROOT / "build" / "Dissertation_FINAL_Uririe_Orume_Dominic.docx"
MD = ROOT / "docs/dissertation/DISSERTATION_FULL.md"

fails, warns = [], []


def ok(cond, msg):
    print(f"  {'ok  ' if cond else 'FAIL'}  {msg}")
    if not cond:
        fails.append(msg)


def warn(cond, msg):
    print(f"  {'ok  ' if cond else 'warn'}  {msg}")
    if not cond:
        warns.append(msg)


def section(t):
    print(f"\n{t}")


# ---------------------------------------------------------------- the document
section("[1] THE SUBMITTED DOCUMENT")
ok(DOCX.exists(), f"{DOCX.name} exists")
doc = Document(DOCX)
z = zipfile.ZipFile(DOCX)
parts = [p.text for p in doc.paragraphs]
for t in doc.tables:
    for r in t.rows:
        parts.extend(c.text for c in r.cells)
full = "\n".join(parts)

imgs = [n for n in z.namelist() if n.startswith("word/media/")]
ok(len(imgs) == 28, f"28 images embedded: the crest, 15 body figures and 12 in Appendix F (found {len(imgs)})")
ok("Appendix F: Supplementary evidence" in full, "Appendix F present")
ev = ROOT / "docs" / "evidence"
ok(len(list(ev.glob("E*.jpg"))) == 12 and (ev / "README.md").exists(),
   "public evidence folder has 12 images and an index")
ok(all(z.getinfo(n).file_size > 4000 for n in imgs), "no figure is a truncated file")
ok(len(doc.tables) == 6, f"6 tables (found {len(doc.tables)})")
for i, t in enumerate(doc.tables, 1):
    filled = sum(1 for r in t.rows for c in r.cells if c.text.strip())
    ok(filled >= len(t.rows) * len(t.columns) * 0.9,
       f"table {i} is fully populated ({filled} of {len(t.rows)*len(t.columns)} cells)")
heads = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
ok(len(heads) >= 55, f"{len(heads)} headings for the outline pane")
ok(sum(1 for p in doc.paragraphs if p.style.name == "Chapter Label") == 6, "6 chapters present")

section("[1b] ASTON TEMPLATE LAYOUT")
from docx.oxml.ns import qn
sec = doc.sections[0]
ok(abs(sec.page_width.cm - 21.0) < 0.05 and abs(sec.page_height.cm - 29.7) < 0.05, "A4 paper")
ok(sec.different_first_page_header_footer, "the cover carries no page number")
pn = sec._sectPr.find(qn("w:pgNumType"))
ok(pn is not None and pn.get(qn("w:start")) == "0", "numbering makes Acknowledgements page 1")
ok(bool(doc.paragraphs[0]._p.xpath(".//pic:pic")), "the crest heads the cover")
cover = " ".join(p.text for p in doc.paragraphs[:16])
for needle in ("Aston University", "Department of AI and Robotics",
               "in fulfilment of the requirements for the degree of",
               "Master of Science in Artificial Intelligence and Business Strategy",
               "Supervisors:", "Word count:"):
    ok(needle in cover, f"cover carries '{needle}'")
front = [p.text for p in doc.paragraphs if p.style.name == "Front Heading"]
ok(front == ["Acknowledgements", "Declaration", "Abstract", "Contents",
             "List of Tables", "List of Figures"], f"front matter in template order {front}")
toc = [p for p in doc.paragraphs if p.style.name.lower().startswith("toc")]
ok(len(toc) >= 60 and all(re.search(r"\d\s*$", p.text) for p in toc),
   f"contents and both lists filled in with page numbers ({len(toc)} entries)")
caps = [p.text for p in doc.paragraphs if re.match(r"^(Figure|Table) [0-9A-Z]+\.\d+:", p.text)]
ok(sum(c.startswith("Figure") for c in caps) == 27, "27 figure captions in the 'Figure 4.1:' form")
ok(sum(c.startswith("Table") for c in caps) == 6, "6 table captions in the 'Table 4.1:' form")
kids = list(doc.element.body.iterchildren())
below = ["".join(kids[i + 1].itertext()).strip() for i, e in enumerate(kids) if e.tag == qn("w:tbl")]
ok(all(b.startswith("Table ") for b in below), "every table caption sits below its table")
apps = [p.text for p in doc.paragraphs if p.style.name == "Heading 2" and p.text.startswith("Appendix ")]
ok(len(apps) == 6, f"appendices A to F are headings ({len(apps)})")
ok("3.8 Ethical considerations" in full and "signed consent" in full, "ethics section with rater consent")
PDF = DOCX.with_suffix(".pdf")
ok(PDF.exists() and PDF.stat().st_mtime >= DOCX.stat().st_mtime - 120, "PDF exported from this build")

section("[1c] EVERY APPENDIX LINK OPENS")
import urllib.request
urls = sorted({r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")})
ok(len(urls) >= 15, f"{len(urls)} distinct links in the document")
dead = []
for u in urls:
    try:
        with urllib.request.urlopen(urllib.request.Request(u, method="HEAD"), timeout=15) as resp:
            if resp.status >= 400:
                dead.append(u)
    except Exception:
        dead.append(u)
ok(not dead, f"all links respond ({dead or 'none dead'})")

section("[2] DOCUMENT HYGIENE")
for label, needle in [
    ("no editorial notices", "delete before submission"),
    ("no author to-dos", "requires you"),
    ("no draft marker", "End of dissertation draft"),
    ("no registry note", "REGISTRY CHECK"),
    ("no em dashes", "—"),
    ("no raw markdown", "!["),
    ("no TODO/FIXME", "TODO"),
    ("no lorem", "lorem ipsum"),
]:
    ok(needle.lower() not in full.lower(), label)
ok("Uririe, Orume Dominic" in full, "author name as verified with the registry")
ok("MSc Artificial Intelligence and Business Strategy" in full, "programme title")
ok("September 2026" in full, "submission month")
aug = [p for p in parts if "August 2026" in p]
ok(all("Accessed" in p or "Git dates every addition" in p for p in aug),
   f"August appears only in citation dates and the dated scope-drift finding "
   f"({len(aug)} mentions)")
ok("I declare that this dissertation is my own work" in full, "Declaration present")
ok("Use of generative AI" in full, "AI-use declaration present")
ok("Ikenna Onyedebelu" in full, "Rater 1 credited by name")
ok("Matthew Brian Tahir" in full, "Rater 2 credited by name")
ok("acted as Rater 1" not in full and "Rater 1 is the author" not in full,
   "the author is no longer described as a rater")
ok("both named here with their consent" in full, "consent recorded for both raters")

# ------------------------------------------------------------ document vs data
section("[3] EVERY HEADLINE FIGURE AGAINST THE DATA")
df = pd.read_csv(ROOT / "data/reports/main_001.csv")
df["cond"] = df.run_id.str.split("__").str[2]


def mean(metric, cond):
    return float(df[(df.metric == metric) & (df.cond == cond)]["value"].mean())


TABLE41 = {
    "hallucinations":   {"claude_code": 0.00, "cursor_agent": 0.17,
                         "replit_agent": 1.33, "antigravity": 0.33},
    "security_density": {"claude_code": 9.65, "cursor_agent": 5.93,
                         "replit_agent": 0.00, "antigravity": 1.47},
    "duplication_pct":  {"claude_code": 0.00, "cursor_agent": 0.90,
                         "replit_agent": 9.56, "antigravity": 4.26},
    "complexity_mean":  {"claude_code": 3.35, "cursor_agent": 2.72,
                         "replit_agent": 2.39, "antigravity": 2.60},
}
for metric, row in TABLE41.items():
    for cond, want in row.items():
        got = mean(metric, cond)
        ok(abs(got - want) < 0.02,
           f"{metric}/{cond}: document {want:.2f}, data {got:.2f}")
        ok(f"{want:.2f}" in full, f"  {want:.2f} appears in the document")

section("[4] KAPPA AS REPORTED")
k = json.loads((ROOT / "data/labels/kappa_results.json").read_text())
for v in ("0.870", "0.852", "0.727"):
    ok(v in full, f"pre-registered kappa {v} quoted")
ok("circular" in full, "circularity of the post-repair value declared")
r1 = {r["item_id"]: r["n_offspec_features"]
      for r in csv.DictReader(open(ROOT / "data/labels/labels_rater1.csv"))}
r2 = {r["item_id"]: r["n_offspec_features"]
      for r in csv.DictReader(open(ROOT / "data/labels/labels_rater2.csv"))}
ok(len(r1) == 19 and len(r2) == 19, "both rater sheets carry 19 items")
ok(all(v != "" for v in r1.values()) and all(v != "" for v in r2.values()),
   "no unlabelled items in either sheet")

section("[5] REPORTS AGREE WITH EACH OTHER")
r = subprocess.run([sys.executable, str(ROOT / "scripts/check_report_consistency.py")],
                   capture_output=True, text=True)
ok(r.returncode == 0, "every report file agrees with Table 4.1")

section("[6] WHAT THE DASHBOARD WOULD SERVE")
ph = pd.read_csv(ROOT / "data/reports/main_001_plus_human.csv")
ph = ph[ph.metric == "security_density"]
for cond, want in [("claude_code", 9.65), ("cursor_agent", 5.93)]:
    got = float(ph[ph.condition == cond]["value"].mean())
    ok(abs(got - want) < 0.02,
       f"deployed report would show {cond} security {want:.2f} (is {got:.2f})")
ok((ROOT / "Dockerfile").read_text().count("COPY data") == 1,
   "Dockerfile copies data into the image")

section("[7] PUBLISHED PACKAGE")
ver = re.search(r'^version = "([^"]+)"', (ROOT / "pyproject.toml").read_text(), re.M)
print(f"        pyproject version: {ver.group(1)}")
deriver = (ROOT / "auditor/analyzers/manifest_deriver.py").read_text()
warn("_JS_ROUTE_RE" in deriver, "route detector fix is present in the source")
dist = sorted((ROOT / "dist").glob(f"*{ver.group(1)}*"))
ok(len(dist) == 2, f"sdist and wheel built for {ver.group(1)} ({len(dist)} found)")
sec = (ROOT / "SECURITY.md").read_text()
ok("No telemetry" in sec, "SECURITY.md states no telemetry")
ok(not (ROOT / "auditor/core/telemetry.py").exists(), "telemetry module removed")
net = [str(p.relative_to(ROOT)) for p in (ROOT / "auditor").rglob("*.py")
       if "urllib.request" in p.read_text(errors="ignore")
       or "smtplib" in p.read_text(errors="ignore") and "dashboard" not in str(p)]
ok(not net, f"no outbound HTTP on any path the CLI reaches ({net or 'none'})")
orphan = [str(p.relative_to(ROOT)) for p in (ROOT / "auditor").rglob("*.pyc")
          if not p.with_suffix("").with_suffix(".py").exists()
          and not (p.parent.parent / (p.stem.split(".")[0] + ".py")).exists()]
ok(not orphan, f"no bytecode left behind by deleted modules ({orphan or 'none'})")

section("[8] WORD COUNT")
lines = MD.read_text().splitlines()
cut = next(i for i, l in enumerate(lines) if re.match(r"^#+\s*References", l, re.I))
body = [l for l in lines[:cut] if not l.startswith(">")]
ch1 = next(i for i, l in enumerate(body) if l.startswith("## 1.1"))
keep, skip = [], False
for l in body[ch1:]:
    if l.startswith("**Figure "):
        skip = True
    elif skip and not l.strip():
        skip = False
    if not skip and not l.startswith("!["):
        keep.append(l)
w = len(re.findall(r"\S+", "\n".join(keep)))
print(f"        chapters excluding figure captions: {w:,}")
ok(w <= 12000, f"under the 12,000 limit ({12000 - w} to spare)")

section("[9] REPOSITORY")
st = subprocess.run(["git", "status", "--porcelain"], cwd=ROOT,
                    capture_output=True, text=True).stdout.strip()
ok(not st, "working tree clean")
t = subprocess.run([sys.executable, "-m", "pytest", "-q"], cwd=ROOT,
                   capture_output=True, text=True).stdout
ok("failed" not in t.split("\n")[-2] if len(t.split("\n")) > 1 else False,
   f"test suite: {t.strip().splitlines()[-1] if t.strip() else 'no output'}")

print(f"\n{'='*60}\n  {len(fails)} failures, {len(warns)} warnings")
for f in fails:
    print(f"   FAIL: {f}")
for f in warns:
    print(f"   warn: {f}")
sys.exit(1 if fails else 0)
