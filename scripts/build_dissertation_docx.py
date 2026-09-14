#!/usr/bin/env python3
"""Render DISSERTATION_FULL.md as the submission document, laid out to the
Aston dissertation template.

Two stages. python-docx writes the .docx: A4, an unnumbered cover carrying the
University crest, front matter in the template's order, two-line chapter
openings, captions below figures and tables, and appendices as real headings.
Microsoft Word then opens that file, fills in the contents and the lists of
tables and figures, and exports the PDF, which is the one file the module
accepts. Only Word knows where the pages break, so page numbers cannot be
computed here.

Editorial blockquotes are dropped, because this artefact is the thing handed
to a marker.

    .venv/bin/python scripts/build_dissertation_docx.py            # .docx and PDF
    .venv/bin/python scripts/build_dissertation_docx.py --no-word  # .docx only
"""
from __future__ import annotations

import platform
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs/dissertation/DISSERTATION_FULL.md"
BASE = SRC.parent
# Deliberately not the name a stale copy already occupies in the author's
# Downloads folder. Three separate reviews were carried out against that
# day-old file because the two were indistinguishable by name.
OUT = ROOT / "build" / "Dissertation_FINAL_Uririe_Orume_Dominic.docx"
PDF = OUT.with_suffix(".pdf")
# Extracted from the Aston template. build/ is gitignored, so the University's
# arms are never published alongside the code.
CREST = ROOT / "build" / "assets" / "aston_university_crest.jpg"

COVER = {
    "institution": "Aston University",
    "department": "Department of AI and Robotics",
    "degree": "Master of Science in Artificial Intelligence and Business Strategy",
    "date": "September 2026",
    "author": "Uririe, Orume Dominic",
    "supervisors": "Julien Barney and Kate Sugden",
    "address": "Aston University, Aston Triangle, Birmingham, B4 7ET, United Kingdom",
    "project": "Project JBKS1",
}

INK = RGBColor(0x1A, 0x1F, 0x1D)
BLACK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x5A, 0x60, 0x5C)

BODY_FONT = "Georgia"
HEAD_FONT = "Georgia"
MONO_FONT = "Consolas"
TEXT_WIDTH_IN = 6.0

# Front matter takes the template's large headings on a fresh page, but is not
# a Heading style: the contents page lists the chapters, not itself.
FRONT = ("Acknowledgements", "Declaration", "Abstract", "Table of Contents")


# ------------------------------------------------------------------ utilities
def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_field(par, instr, result=True, hidden=False):
    """A complex field: begin, instruction, [separate], end. Word computes the
    result when fields are updated."""
    def run_with(el):
        r = par.add_run()
        if hidden:
            r.font.hidden = True
        r._r.append(el)
        return r

    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    run_with(fld("begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    if result:
        run_with(fld("separate"))
    run_with(fld("end"))


def add_tc(par, text, table_id, level=1):
    """A table-of-contents entry. Contents entries carry "C", figures "F",
    tables "T", so each of the three lists collects only its own entries.
    The entry must not be formatted hidden: Word's TOC field skips hidden TC
    fields, and a TC field has no visible result, so nothing shows either way."""
    add_field(par, f'TC "{text.replace(chr(34), chr(39))}" \\f {table_id} \\l {level}',
              result=False)


def add_page_numbers(section):
    """Footer with a live PAGE field, centred, as in the template."""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE")
    for r in p.runs:
        r.font.name = BODY_FONT
        r.font.size = Pt(9.5)


def add_hyperlink(par, url, text, size):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), BODY_FONT)
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "1F4E79")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(round(size * 2)))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    for el in (fonts, colour, sz, u):
        rpr.append(el)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(rpr)
    r.append(t)
    h.append(r)
    par._p.append(h)


INLINE = re.compile(
    r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`|\[[^\]]+?\]\([^)]+?\))", re.S)


def add_runs(par, text, size=11, color=INK, italic_all=False):
    """Render markdown inline emphasis and links into runs."""
    for ch in ("\u00a0", "\u202f", "\u2009"):
        text = text.replace(ch, " ")
    for piece in INLINE.split(text):
        if not piece:
            continue
        bold = italic = mono = False
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            piece, bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            piece, italic = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            piece, mono = piece[1:-1], True
        else:
            m = re.fullmatch(r"\[([^\]]+?)\]\(([^)]+?)\)", piece)
            if m:
                if m.group(2).startswith("http"):
                    add_hyperlink(par, m.group(2), m.group(1), size)
                    continue
                piece = m.group(1)
        run = par.add_run(piece)
        run.bold = bold
        run.italic = italic or italic_all
        run.font.size = Pt(size - (0.6 if mono else 0))
        run.font.name = MONO_FONT if mono else BODY_FONT
        run.font.color.rgb = color


def body_par(doc, text, size=11, space_after=8, justify=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.42
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, text, size=size)
    return p


# --------------------------------------------------------------------- styles
def _plain_style_fonts(style):
    """The default template ties headings to theme fonts and theme colours,
    which override explicit values in Word. Strip the theme attributes so the
    heading is Georgia and black wherever it appears, contents page included."""
    rpr = style.element.rPr
    if rpr is None:
        return
    for el in rpr.iterchildren():
        for attr in list(el.attrib):
            if "Theme" in attr or "theme" in attr:
                del el.attrib[attr]


def ensure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK

    def own(name, size, before, after, page_break):
        s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
        s.font.name = HEAD_FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = BLACK
        pf = s.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.page_break_before = page_break
        pf.keep_with_next = True

    # Without this Word opens the file in Compatibility Mode and lays it out
    # with Word 2010 rules, which is not what a reader's Word will show.
    compat = doc.settings.element.find(qn("w:compat"))
    if compat is not None:
        for cs in compat.findall(qn("w:compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                cs.set(qn("w:val"), "15")

    own("Front Heading", 24, 24, 24, page_break=True)
    own("Chapter Label", 20, 24, 4, page_break=True)
    for level, size, before, after in ((1, 24, 0, 24), (2, 14, 18, 8),
                                       (3, 11.5, 14, 6)):
        h = styles[f"Heading {level}"]
        _plain_style_fonts(h)
        h.font.name = HEAD_FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = BLACK
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True


HEAD_SIZES = {1: 24, 2: 14, 3: 11.5}


def heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    add_runs(p, text, size=HEAD_SIZES[min(level, 3)], color=BLACK)
    for r in p.runs:
        r.bold = True
        r.font.name = HEAD_FONT
    return p


def chapter_heading(doc, number, title):
    """Two lines, as in the template: "Chapter 1", then "Introduction"."""
    label = doc.add_paragraph(style="Chapter Label")
    label.add_run(f"Chapter {number}")
    h = heading(doc, title, 1)
    add_tc(h, f"{number}  {title}", "C")


def front_heading(doc, text):
    p = doc.add_paragraph(style="Front Heading")
    p.add_run(text)
    return p


def add_toc(doc, instr):
    p = doc.add_paragraph()
    add_field(p, instr)


# -------------------------------------------------------------------- figures
# Figures are generated at 300 dpi for print. Embedding them at full size
# makes the .docx too large to transport; 1150 px across a 6-inch column is
# still 190 dpi, which survives PDF export without visible loss.
MAX_PX = 1150
_TMP = ROOT / "build" / "_figures_optimised"


def optimised(path: Path) -> Path:
    _TMP.mkdir(parents=True, exist_ok=True)
    out = _TMP / path.name
    if out.exists() and out.stat().st_mtime >= path.stat().st_mtime:
        return out
    with Image.open(path) as im:
        im = im.convert("RGB")
        if im.width > MAX_PX:
            h = round(im.height * MAX_PX / im.width)
            im = im.resize((MAX_PX, h), Image.LANCZOS)
        # These are charts and line diagrams: a 256-colour palette is visually
        # lossless here and roughly a third the size of truecolour PNG.
        im.quantize(colors=256, method=Image.MEDIANCUT).save(
            out, "PNG", optimize=True)
    return out


CAP_RE = re.compile(r"^\*\*(Figure|Table) ([0-9A-Z]+\.[0-9]+)\*\*")


def caption(doc, text, titles, space_before=4):
    """Caption with the template's colon ("Figure 4.1:") and a hidden entry
    that places it in the list of figures or tables."""
    m = CAP_RE.match(text)
    kind, num = m.group(1), m.group(2)
    if (kind, num) not in titles:
        raise SystemExit(f"refusing to build: {kind} {num} is captioned but "
                         f"missing from the List of {kind}s")
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = cp.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(14)
    pf.left_indent = Inches(0.25)
    pf.right_indent = Inches(0.25)
    pf.line_spacing = 1.15
    add_tc(cp, f"{num}  {titles[(kind, num)]}", kind[0])
    add_runs(cp, CAP_RE.sub(f"**{kind} {num}:**", text, count=1), size=9.5)
    return cp


def add_figure(doc, path: Path, cap: str, titles):
    path = optimised(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(TEXT_WIDTH_IN))
    if cap:
        caption(doc, cap, titles)


# --------------------------------------------------------------------- tables
def _borders(table, colour="C9CFCB"):
    """Hairline grid; the data should carry the emphasis, not the rules."""
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), colour)
        borders.append(el)
    pr.append(borders)


NUMERIC = re.compile(r"^[−\-]?[\d.,]+\s*%?$|^[\d.]+\s*×\s*10", re.U)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    t = doc.add_table(rows=1, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    _borders(t)

    for i, cell in enumerate(t.rows[0].cells):
        cell.text = ""
        set_cell_bg(cell, "0F514B")
        par = cell.paragraphs[0]
        par.paragraph_format.space_before = Pt(3)
        par.paragraph_format.space_after = Pt(3)
        par.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if i and i == ncols - 1
                         else WD_ALIGN_PARAGRAPH.LEFT)
        add_runs(par, header[i].strip(), size=9.2,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
        for r in par.runs:
            r.bold = True

    for k, row in enumerate(body):
        cells = t.add_row().cells
        if k % 2 == 1:
            for c in cells:
                set_cell_bg(c, "F4F6F4")
        for i, val in enumerate(row[:ncols]):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_before = Pt(2.5)
            par.paragraph_format.space_after = Pt(2.5)
            v = val.strip()
            if i and NUMERIC.match(re.sub(r"[*`]", "", v)):
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_runs(par, v, size=9.2)
    # the caption now sits below the table; keep the two on one page
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.keep_with_next = True
    return t


# ---------------------------------------------------------------- cover page
def cover(doc, title, words):
    """The template's cover: crest, institution, department, title, the
    fulfilment statement, date, author, supervision and word count."""
    if not CREST.exists():
        raise SystemExit(f"refusing to build: the Aston crest is missing at {CREST}")
    crest = _TMP / "aston_university_crest.jpg"
    _TMP.mkdir(parents=True, exist_ok=True)
    if not crest.exists() or crest.stat().st_mtime < CREST.stat().st_mtime:
        with Image.open(CREST) as im:
            im.convert("RGB").save(crest, "JPEG", quality=90, optimize=True)

    def line(text="", size=11, bold=False, after=4):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.2
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)
            r.font.name = BODY_FONT
            r.font.color.rgb = BLACK
        return p

    p = line(after=6)
    p.add_run().add_picture(str(crest), width=Cm(4.8))
    line(COVER["institution"], 16, after=2)
    line(COVER["department"], 11, after=28)
    line(title, 16, bold=True, after=30)
    line("A dissertation submitted in fulfilment of the requirements for the degree of",
         11, after=0)
    line(COVER["degree"] + ".", 11, after=4)
    line(COVER["date"], 11, after=26)
    line(COVER["author"], 13, bold=True, after=26)
    p = line(after=0)
    for text, bold in (("Supervisors: ", True), (COVER["supervisors"], False)):
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(11)
        r.font.name = BODY_FONT
        r.font.color.rgb = BLACK
    line(COVER["address"], 10, after=4)
    line(COVER["project"], 10, after=26)
    line(f"Word count: {words:,} (main text, excluding figure captions, "
         "references and appendices)", 10, after=0)


# ------------------------------------------------------------------- parsing
FIG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")


def list_titles(lines):
    """Short titles for the lists of tables and figures, taken from the
    manuscript's own lists so the two cannot drift apart."""
    start = next(i for i, l in enumerate(lines) if l.strip() == "## Table of Contents")
    end = next(i for i in range(start, len(lines)) if lines[i].startswith("# Chapter"))
    titles, cur = {}, None
    for l in lines[start:end]:
        m = re.match(r"^- (Figure|Table) ([0-9A-Z]+\.[0-9]+) (.*)$", l)
        if m:
            cur = (m.group(1), m.group(2))
            titles[cur] = m.group(3).strip()
        elif cur and l.startswith("  ") and l.strip():
            titles[cur] += " " + l.strip()
        else:
            cur = None
    return {k: re.sub(r"[*`]", "", v) for k, v in titles.items()}


def render(doc, lines, titles):
    # The manuscript opens with a metadata block and editorial notes. The
    # cover already carries that information; rendering it again produced a
    # second, repeated cover page.
    i = next(k for k, l in enumerate(lines) if l.startswith("## "))
    n = len(lines)
    while i < n:
        line = lines[i]

        # editorial notes -- never reach the marker
        if line.startswith(">") or not line.strip() or line.strip() == "---":
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            i += 1
            if level == 1:
                cm = re.match(r"^Chapter (\d+):\s*(.*)$", text)
                if cm:
                    chapter_heading(doc, cm.group(1), cm.group(2))
                else:
                    h = heading(doc, text, 1)
                    h.paragraph_format.page_break_before = True
                    add_tc(h, text, "C")
            elif text == "Table of Contents":
                front_heading(doc, "Contents")
                # built wholly from entries, so chapters read "1  Introduction"
                # as in the template while the page shows two lines
                add_toc(doc, r"TOC \f C \h \z")
                front_heading(doc, "List of Tables")
                add_toc(doc, r"TOC \f T \h \z")
                front_heading(doc, "List of Figures")
                add_toc(doc, r"TOC \f F \h \z")
                # the manuscript's static lists are replaced by the fields
                i = next(k for k in range(i, n) if lines[k].startswith("# Chapter"))
            elif text in FRONT:
                front_heading(doc, text)
            else:
                add_tc(heading(doc, text, level), text, "C", min(level, 3))
            continue

        fm = FIG_RE.match(line)
        if fm:
            path = (BASE / fm.group(2)).resolve()
            j = i + 1
            while j < n and not lines[j].strip():
                j += 1
            cap = ""
            if j < n and CAP_RE.match(lines[j]):
                cap_lines = []
                while j < n and lines[j].strip():
                    cap_lines.append(lines[j].strip())
                    j += 1
                cap = " ".join(cap_lines)
                i = j
            else:
                i += 1
            if not path.exists():
                raise SystemExit(f"refusing to build: figure missing at {path}")
            add_figure(doc, path, cap, titles)
            continue

        # A table caption precedes its table in the manuscript; the template
        # sets it below, so it is held until the table has been drawn.
        if CAP_RE.match(line):
            cap = [line.strip()]
            i += 1
            while i < n and lines[i].strip() and not lines[i].lstrip().startswith("|"):
                cap.append(lines[i].strip())
                i += 1
            j = i
            while j < n and not lines[j].strip():
                j += 1
            text = " ".join(cap)
            if j < n and lines[j].lstrip().startswith("|"):
                rows = []
                while j < n and lines[j].lstrip().startswith("|"):
                    cells = [c for c in lines[j].strip().strip("|").split("|")]
                    if not re.fullmatch(r"[\s:\-]+", "".join(cells)):
                        rows.append(cells)
                    j += 1
                add_table(doc, rows)
                caption(doc, text, titles, space_before=6)
                i = j
            else:
                caption(doc, text, titles, space_before=10)
            continue

        if line.lstrip().startswith("|"):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not re.fullmatch(r"[\s:\-]+", "".join(cells)):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
                doc.add_paragraph().paragraph_format.space_after = Pt(10)
            continue

        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if lm:
            indent = len(lm.group(1)) // 2
            ordered = bool(re.match(r"\d+\.", lm.group(2)))
            text = [lm.group(3)]
            i += 1
            while i < n and lines[i].strip() and not re.match(
                    r"^(\s*)([-*]|\d+\.)\s+", lines[i]) and not lines[i].startswith("#"):
                text.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph(
                style="List Number" if ordered else "List Bullet")
            p.paragraph_format.left_indent = Inches(0.3 + 0.25 * indent)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.3
            add_runs(p, " ".join(text), size=10.6)
            continue

        # ordinary paragraph
        buf = []
        while i < n and lines[i].strip() and not lines[i].startswith(("#", "|", ">", "!")) \
                and not re.match(r"^(\s*)([-*]|\d+\.)\s+", lines[i]) \
                and not CAP_RE.match(lines[i]) and lines[i].strip() != "---":
            buf.append(lines[i].strip())
            i += 1
        if buf:
            body_par(doc, " ".join(buf))
        else:
            i += 1


# ------------------------------------------------------------------ the Word pass
def _osa(script, timeout):
    r = subprocess.run(["osascript", "-e", script], capture_output=True,
                       text=True, timeout=timeout)
    if r.returncode:
        raise SystemExit(f"Word pass failed: {r.stderr.strip()}")
    return r.stdout.strip()


# Word is sandboxed: it can open a file it is handed, but it can only write
# inside its own container without asking the user. Writing the PDF anywhere
# else stalls on a permission prompt, so both outputs are saved there and
# copied out.
WORD_DATA = Path.home() / "Library/Containers/com.microsoft.Word/Data/Documents"


def finalise_in_word():
    """Fill in the contents and both lists, save, and export the PDF.

    Word's scripting interface on this platform cannot close a document, so
    each pass uses a name of its own: reopening a name Word still holds would
    raise a dialog and stall. The working copy stays open in Word afterwards."""
    if platform.system() != "Darwin" or not Path("/Applications/Microsoft Word.app").exists():
        raise SystemExit("Word for Mac is needed to fill in the contents and export "
                         "the PDF; rerun with --no-word to build the .docx alone")
    WORD_DATA.mkdir(parents=True, exist_ok=True)
    for old in WORD_DATA.glob("dissertation_word_pass_*"):
        old.unlink()
    stem = "dissertation_word_pass_" + datetime.now().strftime("%Y%m%d-%H%M%S")
    name = stem + ".docx"
    # Word silently ignores a file handed to it by AppleScript's "open", even
    # from its own container, but accepts one opened through Launch Services
    # as Finder does. It loads asynchronously, so the script waits for it.
    saved_docx, saved_pdf = WORD_DATA / name, WORD_DATA / (stem + ".pdf")
    shutil.copy2(OUT, saved_docx)
    subprocess.run(["open", "-a", "Microsoft Word", str(saved_docx)], check=True)
    steps = [
        ("open", 'repeat 240 times\n'
                 f'if (name of every document) contains "{name}" then return "{name}"\n'
                 'delay 0.5\nend repeat\nerror "Word did not open the document"'),
        ("fill in contents and lists",
         f'set d to document "{name}"\n'
         'set n to count of tables of contents of d\n'
         'repeat with i from 1 to n\nupdate (table of contents i of d)\nend repeat\n'
         'return n'),
        ("save .docx", f'save as document "{name}" file name "{saved_docx}" '
                       'file format format document\nreturn "saved"'),
        ("export PDF", f'save as document "{name}" file name "{saved_pdf}" '
                       'file format format PDF\nreturn "exported"'),
    ]
    for label, body in steps:
        t0 = time.time()
        script = ('with timeout of 600 seconds\ntell application "Microsoft Word"\n'
                  f'{body}\nend tell\nend timeout')
        out = _osa(script, timeout=660)
        print(f"  word: {label} ({time.time() - t0:.0f}s) {out}")
    if not (saved_docx.exists() and saved_pdf.exists()):
        raise SystemExit("Word pass failed: Word did not write both files")
    shutil.copy2(saved_docx, OUT)
    shutil.copy2(saved_pdf, PDF)
    print(f"  word: the working copy {name} is still open in Word; close it "
          "without saving when convenient")


def main():
    use_word = "--no-word" not in sys.argv[1:]
    text = SRC.read_text()
    lines = text.splitlines()

    # word count for the cover (chapters, excluding figure captions)
    cut = next(k for k, l in enumerate(lines) if re.match(r"^#+\s*References", l, re.I))
    body = [l for l in lines[:cut] if not l.startswith(">")]
    ch1 = next(k for k, l in enumerate(body) if l.startswith("## 1.1"))
    keep, skip = [], False
    for l in body[ch1:]:
        if l.startswith("**Figure "):
            skip = True
        elif skip and not l.strip():
            skip = False
        if not skip and not l.startswith("!["):
            keep.append(l)
    words = len(re.findall(r"\S+", "\n".join(keep)))
    if words > 12000:
        raise SystemExit(f"refusing to build: {words:,} words exceeds the 12,000 limit")

    doc = Document()
    sec = doc.sections[0]
    # A4 with 2.54 cm margins, as in the template
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.54)
    # The cover carries no number and Acknowledgements is page 1.
    sec.different_first_page_header_footer = True
    sec.first_page_footer.is_linked_to_previous = False
    start = OxmlElement("w:pgNumType")
    start.set(qn("w:start"), "0")
    sec._sectPr.insert_element_before(
        start, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
        "w:printerSettings", "w:sectPrChange")
    add_page_numbers(sec)
    ensure_styles(doc)

    title = lines[0].lstrip("# ").strip()
    cover(doc, title, words)
    render(doc, lines[1:], list_titles(lines))

    # Refuse to ship a document containing text written for the author.
    # A leak here reaches a marker, so it is a build failure, not a warning.
    parts = [par.text for par in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.extend(cell.text for cell in row.cells)
    rendered = "\n".join(parts)
    if "\u2014" in rendered:
        raise SystemExit("refusing to build: em dashes present in the text")
    banned = ["delete before submission", "End of dissertation draft",
              "Remaining before submission", "requires you", "REGISTRY CHECK",
              "Editorial status", "TODO", "FIXME", "Reproduce in full",
              "Right-click", "Update field"]
    leaks = [b for b in banned if b.lower() in rendered.lower()]
    if leaks:
        raise SystemExit(f"refusing to build: editorial text leaked -> {leaks}")

    # Stamp the build into the file's own properties. A reader who is unsure
    # whether they have the current version can check File > Properties instead
    # of guessing from a filename, which is how a day-old copy in a Downloads
    # folder came to be reviewed as though it were the submission.
    try:
        commit = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT,
                                capture_output=True, text=True).stdout.strip()
    except Exception:
        commit = "unknown"
    built = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    core = doc.core_properties
    core.title = title
    core.author = COVER["author"]
    core.subject = "MSc Artificial Intelligence and Business Strategy, Aston University"
    core.category = "Dissertation, project JBKS1"
    core.keywords = ("agentic AI; code generation; software quality metrics; "
                     "specification fidelity; AI governance")
    core.comments = (f"Built {built} from commit {commit}. "
                     f"{words:,} words of main text excluding figure captions. "
                     "Generated by scripts/build_dissertation_docx.py; edit "
                     "docs/dissertation/DISSERTATION_FULL.md and rebuild.")
    core.last_modified_by = COVER["author"]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"  build stamp: {built}, commit {commit}")
    print(f"  cover word count: {words:,}")
    if use_word:
        finalise_in_word()
        print(f"wrote {PDF}  ({PDF.stat().st_size/1024:.0f} KB)")
    print(f"wrote {OUT}  ({OUT.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
